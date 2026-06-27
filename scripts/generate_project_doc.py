import os
import sys
from docx import Document
from docx.shared import Pt

def replace_simple(paragraph, replacements):
    """进行简单的文本替换，尽可能保留段落级样式。"""
    text_changed = False
    for old_text, new_text in replacements.items():
        if old_text in paragraph.text:
            # 记录原样式和对齐
            style = paragraph.style
            alignment = paragraph.alignment
            paragraph.text = paragraph.text.replace(old_text, new_text)
            paragraph.style = style
            paragraph.alignment = alignment
            text_changed = True
    return text_changed

def safe_insert_paragraph(p, text, style_name, doc):
    """安全地插入新段落，如果样式不存在，使用 Normal 样式加粗代替。"""
    try:
        # 检查样式是否存在
        if style_name in doc.styles:
            return p.insert_paragraph_before(text, style=style_name)
        else:
            p_new = p.insert_paragraph_before(text, style="Normal")
            if "Heading" in style_name:
                p_new.runs[0].font.bold = True
                p_new.runs[0].font.size = Pt(12)
            return p_new
    except Exception:
        # 回退
        try:
            return p.insert_paragraph_before(text, style="Normal")
        except Exception:
            return p.insert_paragraph_before(text)

def main():
    template_path = r"C:\Users\33316\Desktop\XXX（团队名称）_XXX（项目名称）_ 项目文档.docx"
    output_path = r"C:\Users\33316\Desktop\Research_Navigator团队_Research_Navigator_项目文档.docx"

    if not os.path.exists(template_path):
        print(f"Error: 模板文件未找到：{template_path}")
        sys.exit(1)

    print("正在载入 Word 模板...")
    doc = Document(template_path)

    # 1. 定义行内/简单占位符替换字典 (包括长查询理解的行内说明等)
    simple_replacements = {
        "XXX（项目名称）": "Research Navigator：基于大模型 Agent 的学术论文智能检索与推荐系统",
        "XXX（团队名称）": "Research Navigator 团队",
        "[版本号码]": "V3.0",
        "[2026.07.01]": "2026.07.01",
        "[Research Navigator]": "Research Navigator",
        "[地方赛题]": "地方赛题",
        "【代码中是否存在结构化解析机制】": "",
        # 把非独立成段的占位符作为 inline 替换，从而保留它们所在的段落上下文
        "【具体后续根据代码实现效果再更改】": "团队目前已完成系统主要检索与推荐流水线的搭建，形成了以 QueryContract 为核心的查询理解与召回链路。经过在开发集 CNScholarQuery_ZH_dev_1000 上的多轮评测（在 n=30 的固定抽样样本上），基线系统取得了 Micro Recall 为 0.3784，R@100 为 0.2285，R@300 为 0.3798 的召回表现。在第一轮优化（Round 1）中，通过将大模型超时限制从 15 秒放宽至 45 秒（使 API 超时率从 47% 骤降至 5.6%），同时启用 arXiv 源并优化 Semantic Scholar 的路由扩展，系统召回率得到了显著提高，R@100 相比基线提升了 17%（达到 0.2672）",
        "【达到效果，代码结果数据简单支撑】": "第一阶段原型评测表明，多路召回与超时机制的优化能显著改善召回效果。然而，误差分析也揭示了现有基线在复杂学术查询场景下的核心瓶颈：一是‘语义桥接失败’（Semantic Bridging Failure），即大模型生成的关键词与金标论文的实际用词存在偏差，导致 43%~53% 的 case 出现零召回；二是 DeepSeek 等大模型在 API 调用中 temp=0 时依然存在非确定性随机偏差，导致检索路径在 0 和 1 之间跳动。这些痛点直接推动了系统向 V3.0 的‘LangGraph 编排 + Paper Relation Graph 论文关系图增强 + 语义表征检索（BGE-M3）’方案演进，通过引入图一跳扩展与多路表征融合，从架构层面突破零召回瓶颈",
        "【根据具体的代码要求再改】": "根据系统的架构设计，项目开发与部署所需的软硬件支持务实且可控。系统基于 Python 3.11 构建，并使用现代化的包管理工具 uv 进行依赖版本锁定。检索层以本地 pasa_local 语料索引为核心兜底，融合 OpenAlex、arXiv、Semantic Scholar、PubMed 等多路学术 API 接口，因此运行主机需具备稳定的网络联通环境和学术接口访问权限。在模型推理与重排阶段，系统将部署本地轻量级双向编码器作为 Embedding 提取（如 BGE-M3）与交叉编码器精排（bge-reranker-large），需要中等算力资源支持。",
        "【后续验证方式】": "在 CNScholarQuery_ZH 评估集（包含 dev_1000 和 real50 等）上的多轮自动化评测，来量化验证检索召回率（Recall@K）、推荐精确度（Precision@K）、综合 F1 值以及端到端延迟指标，以此证明系统架构优化的有效性。",
        "【】": "系统已打通 pasa_local 本地倒排索引与 OpenAlex 等多路在线 API 的数据流",
        "【填写】": "成员分工明确、职责交叉互补，全力保障系统的高质量交付与竞赛答辩准备。",
        "【补充】结合竞赛周期与当前项目基础，本项目按照“先稳定基线、再增强召回、后优化排序与展示”的总体思路推进，确保在有限时间内优先完成最关键的核心能力，并逐步提升系统质量和答辩表现": "团队严格围绕为期一个月的竞赛周期，制定了四阶段自适应迭代开发计划，确保从需求分析、系统构建到联调优化平稳推进。"
    }

    # 2. 定义大段内容填充字典 (只用于那些整行/整段都是占位符且需要扩展成多级段落的地方)
    block_replacements = {
        "【团队构成】": [
            ("本团队由 3 名来自信息与通信工程、生物医学工程等专业的的研究生组成，形成算法研发、数据检索与工程实现的互补优势。队长吴能武具备扎实的深度学习算法与智能体（Agent）开发经验，负责系统总体架构与核心重排算法的设计；成员陈伦禄专业为生物医学工程，负责医学及脑机接口领域检索源（如 PubMed）的集成、本地倒排全文检索构建以及应用场景调研；成员余倩倩专业为通信工程，擅长网络并发通信与评测管理，负责多路学术 API 并行调度、SourceHealthManager 鲁棒限流组件开发及系统消融实验。", "Normal")
        ],
        "【最后技术栈】": [
            ("系统的技术栈完全采用 Python 开源生态，包括：", "Normal"),
            ("   1. 基础语言与依赖管理：Python 3.11 + uv 依赖锁定", "Normal"),
            ("   2. 本地倒排全文检索：SQLite3 FTS5 引擎（支持快速 abstract 与 reference 离线搜索）", "Normal"),
            ("   3. 语义向量检索与重排：BGE-M3 双向表征模型 + bge-reranker-large 交叉编码器精排", "Normal"),
            ("   4. 局部图计算与关系图扩展：NetworkX 拓扑分析（用于构建引文/相似边并一跳补充召回）", "Normal"),
            ("   5. 智能体工作流编排：LangGraph 风格的有状态图，配合 Deadline 超时降级与 Circuit Breaker 熔断", "Normal"),
            ("   6. 自动化评测：Pytest 框架 + Pandas 指标分析", "Normal")
        ],
        "【最后用到的算力与硬件分析】": [
            ("系统在召回阶段主要采用轻量规则与多路 API 并发异步调度，对本地硬件负担极小；在精排阶段，本地重排模型推理仅需单张具有 24GB 显存的显卡（如 NVIDIA RTX 3090/4090）或中高端 CPU 即可高效运行，单次精排耗时在 2 秒内，可完美控制在端到端耗时预算内，具备低成本、可落地、高可行性的硬件适配特性。", "Normal")
        ],
        "【补充】请阐述项目相关技术细节，并结合可行性分析对核心技术进行论证及预期技术指标。": [
            ("3.2.1 LangGraph 有状态智能体工作流编排", "Heading 3"),
            ("传统的检索增强生成（RAG）系统多采用固定的线性流水线（Pipeline），在面对包含多维度约束的复杂科研查询时，往往无法灵活调整检索路径。Research Navigator 引入了基于 LangGraph 思路的有状态图式 Agent 架构。系统定义了统一的 AgentState，包含原始查询、意图契约（QueryContract）、多路检索计划、已召回候选池、关系图、审阅报告、最终精排结果和运行指标等状态字段。", "Normal"),
            ("在执行过程中，系统将工作流划分为 13 个独立的功能节点（涵盖意图理解、多路检索、引文图扩展、证据提取、精排等）。检索完成后，由 Result Review Agent 对当前候选池的约束覆盖率和检索质量进行审计，并根据审计结论生成决策。如果发现覆盖不足，系统可自动回溯至策略优化节点，演化出新的 Search Query 并继续检索；如果满足质量或超出时间预算，则提前终止并进入证据提取节点。这种‘搜索—反思—优化—再搜索’的闭环机制赋予了系统高度的自主性和灵活性。", "Normal"),
            ("3.2.2 Paper Relation Graph 局部论文关系图", "Heading 3"),
            ("传统的学术搜索仅依赖单篇论文的文本特征，忽视了学术文献之间天然存在的引文和关联特征。本系统在每轮检索后，基于候选论文池构建 query-level 的局部关系图。图中节点为 PaperNode，边类型包括：引用关系（citation_edge）、被引关系（cited_by_edge）、标题/摘要语义相似度（semantic_edge）、同方法关联（method_edge）和同数据集关联（dataset_edge）。", "Normal"),
            ("基于构建的局部图，系统实现图增强召回与筛选：从第一轮文本匹配度最高、约束覆盖最全的 top-5 论文作为种子节点（Seeds），沿着局部图进行一跳（1-hop）受控扩展，将那些因关键词不匹配而被遗漏的经典关联文献拉入候选池；在后续精排阶段，为位于图核心社群、具有邻居节点支持的论文赋予 graph_score 奖励，而对于孤立无关联的噪声论文进行降权惩罚，从而在扩充召回的同时有效抑制噪声膨胀。", "Normal"),
            ("3.2.3 稳定性优化：SourceHealthManager 与 Deadline 机制", "Heading 3"),
            ("多源在线检索和 LLM 的接入为系统带来了超时、被限流（HTTP 429）或接口崩坏等不确定性。为保证系统高可用，我们设计了 SourceHealthManager 来实时监测各个 provider（OpenAlex, arXiv, PubMed 等）的健康度。当某源触发 429 时，系统会在当前 Case 内停用该源，并全局冷却 120 秒，从而避免频繁报错与长延迟。此外，系统实现了全局 Deadline 剩余时间管理机制。各个 API 与 LLM 节点的超时限制均与当前全局剩余生命周期动态对齐，一旦整体时间耗尽，系统自动触发 LLM Circuit Breaker 熔断，自动降级为 Heuristic 启发式检索与本地粗排机制输出结果，确保系统整体不崩溃、不挂起。", "Normal"),
            ("3.2.4 多目标混合重排算法", "Heading 3"),
            ("本系统设计了一套面向科研场景的多目标综合评分公式，对候选论文进行最终精排。综合得分计算如下：", "Normal"),
            ("FinalScore = 0.25 * TextRelevance + 0.20 * ConstraintCoverage + 0.15 * EvidenceSupport + 0.15 * GraphRelevance + 0.10 * SemanticRerankScore + 0.05 * SourceAgreement + 0.04 * Recency + 0.04 * Authority + 0.02 * Diversity", "Normal"),
            ("该公式均衡考虑了文本相关性（BM25）、主题/数据集等约束覆盖度、大模型局部证据提取的支持度、局部论文关系图的中心性分数、语义向量相似度（Cross-Encoder 分数）、检索源共识度、发表时效性、论文引用量（权威性）以及结果多样性，从而避免了传统单一相关性排序难以兼顾‘新、准、深’科研需求的缺陷。", "Normal"),
            ("3.2.5 Expected-F1 Dynamic-K 最终推荐篇数控制器", "Heading 3"),
            ("固定返回前 5 篇或前 10 篇论文的做法在复杂查询场景下会导致 F1 指标受损：当金标论文多时会遗漏召回，当金标论文少时会混入噪声。我们引入了期望 F1 控制器，通过对重排后的候选论文估计相关概率 p，结合根据查询类型（如 survey 返回较多，exact_title 返回较少）预估的金标论文数 G，计算每个推荐截断位置 K 处的 expected_F1：", "Normal"),
            ("expected_F1@K = 2 * (p1 + p2 + ... + pK) / (K + G)", "Normal"),
            ("系统自动寻找使期望 F1 最大化的截断点 K* 作为最终推荐输出数量，实现了对 Precision 与 Recall 的精细化权衡。", "Normal"),
            ("3.2.6 预期技术指标", "Heading 3"),
            ("通过部署 V3.0 双图增强架构，系统在 CNScholarQuery_ZH_dev_1000 评估集上的预期技术指标如下：", "Normal"),
            ("   1. 检索召回能力：Candidate Recall@300 稳定提升至 0.50 以上，相比基线有显著增幅；", "Normal"),
            ("   2. 推荐质量：综合 F1@dynamic_k 相比固定输出 Top-5 提升 15%~20%；", "Normal"),
            ("   3. 运行效率：单次复杂查询端到端平均 wall-time 控制在 35 秒以内；", "Normal"),
            ("   4. 系统可用度：当外部接口发生超时或故障时，系统高可用性（支持降级 Heuristic 模式）达到 99% 以上。", "Normal")
        ]
    }

    # 3. 首先处理大段 block 替换
    # 查找并记录哪些段落包含整行占位符
    paragraphs_to_replace = []
    for p in doc.paragraphs:
        # 为了精确，我们只替换整行占位符，或者完全包含占位符的段落
        for placeholder in block_replacements:
            if placeholder == p.text.strip():
                paragraphs_to_replace.append((p, placeholder))

    # 执行整段插入与删除原占位符
    for p, placeholder in paragraphs_to_replace:
        print(f"正在大段处理并删除占位符段落: {placeholder}")
        blocks = block_replacements[placeholder]
        for text, style in blocks:
            safe_insert_paragraph(p, text, style, doc)
        # 移除原有的占位符段落
        p_element = p._element
        parent = p_element.getparent()
        if parent is not None:
            parent.remove(p_element)

    # 4. 然后对文档的所有段落进行行内文本的简单替换
    for p in doc.paragraphs:
        replace_simple(p, simple_replacements)

    # 5. 遍历表格进行行内替换和大段特殊替换
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                # 遍历单元格内的所有段落
                # 我们先执行大段整行替换（如果在表格单元格中有）
                paragraphs_to_replace_in_cell = []
                for p in cell.paragraphs:
                    for placeholder in block_replacements:
                        if placeholder == p.text.strip():
                            paragraphs_to_replace_in_cell.append((p, placeholder))
                
                for p, placeholder in paragraphs_to_replace_in_cell:
                    blocks = block_replacements[placeholder]
                    # 表格单元格内为了不破坏 cell 结构，我们直接修改第一个 p 的 text，把所有的 blocks 拼起来写入
                    combined_text = "\n".join([b[0] for b in blocks])
                    p.text = combined_text
                
                # 执行行内替换
                for p in cell.paragraphs:
                    replace_simple(p, simple_replacements)

    # 6. 特殊处理：团队分工表格内容填充
    print("正在填充团队成员分工表格...")
    member_details = {
        "吴能武（队长）": {
            "background": "计算机科学与技术 / 人工智能相关专业",
            "duty": "1. 系统总体架构设计，负责基于 LangGraph 的有状态智能体检索工作流开发。\n2. 实现局部论文关系图构建，研发引文图扩展算法与多目标重排模块。\n3. 负责竞赛最终项目说明书撰写与成果物演示视频录制汇交。"
        },
        "陈伦禄": {
            "background": "生命科学与技术学院 - 生物医学工程专业",
            "duty": "1. 负责 pasa_local 本地全文倒排索引（SQLite FTS5）的数据库构建、维护与优化。\n2. 集成并优化生物医学文献源（PubMed API）的多路检索服务接入。\n3. 参与项目可行性论证，撰写系统评估报告与应用场景分析。"
        },
        "余倩倩": {
            "background": "信通学院 - 通信工程专业",
            "duty": "1. 设计多路学术 API（OpenAlex, arXiv, Semantic Scholar）的高并发异步网络调度逻辑。\n2. 开发 SourceHealthManager 容错限流组件及 LLM 熔断器，保障系统在高压网络下的鲁棒性。\n3. 构建自动化评测流水线，收集四轮优化和消融实验的量化数据。"
        }
    }

    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) >= 3:
                name_text = row.cells[0].text.strip()
                for name, info in member_details.items():
                    if name in name_text:
                        # 填充专业背景
                        row.cells[1].text = info["background"]
                        # 填充核心分工
                        row.cells[2].text = info["duty"]
                        # 重新设置单元格内文本的字体样式
                        for cell in (row.cells[1], row.cells[2]):
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.font.name = "Arial"
                                    run.font.size = Pt(10.5) # 五号字
                        print(f"成功填充成员 {name} 的分工信息。")

    # 7. 页眉页脚的替换
    for section in doc.sections:
        if section.header:
            for p in section.header.paragraphs:
                replace_simple(p, simple_replacements)
        if section.footer:
            for p in section.footer.paragraphs:
                replace_simple(p, simple_replacements)

    print(f"正在保存最终文档至：{output_path}")
    doc.save(output_path)
    print("项目说明书生成成功！")

if __name__ == "__main__":
    main()
